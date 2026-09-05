import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell internal padding."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """Sets subtle borders for a table."""
    tbl_pr = table._tbl.tblPr
    tbl_borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(tbl_borders)

def build_word_document():
    doc = Document()
    
    # Page Setup - Margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    # Color Palette
    PRIMARY_COLOR = RGBColor(27, 54, 93)     # Navy #1B365D
    SECONDARY_COLOR = RGBColor(44, 94, 138)  # Slate Blue #2C5E8A
    BODY_COLOR = RGBColor(40, 40, 40)        # Charcoal #282828
    MUTED_COLOR = RGBColor(100, 100, 100)    # Gray #646464
    
    # Base Style configuration
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = BODY_COLOR
    
    # ----------------------------------------------------
    # DOCUMENT HEADER / BANNER
    # ----------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run("Assignment 5: Social Network Analysis with R")
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = PRIMARY_COLOR
    
    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(0)
    subtitle_p.paragraph_format.space_after = Pt(12)
    sub_run = subtitle_p.add_run("R Programming Laboratory Report | igraph Network Analysis & Visualization")
    sub_run.font.size = Pt(12)
    sub_run.font.italic = True
    sub_run.font.color.rgb = SECONDARY_COLOR
    
    # Info Box Table
    info_table = doc.add_table(rows=2, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    
    col_widths = [Inches(3.3), Inches(3.3)]
    for row in info_table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]
            set_cell_background(cell, "F2F5F8")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            
    info_table.rows[0].cells[0].paragraphs[0].add_run("Course: ").bold = True
    info_table.rows[0].cells[0].paragraphs[0].add_run("R Programming Lab")
    info_table.rows[0].cells[1].paragraphs[0].add_run("Dataset: ").bold = True
    info_table.rows[0].cells[1].paragraphs[0].add_run("networkdata.csv (52 Nodes, 290 Edges)")
    
    info_table.rows[1].cells[0].paragraphs[0].add_run("Topic: ").bold = True
    info_table.rows[1].cells[0].paragraphs[0].add_run("Graph Theory & Social Network Analytics")
    info_table.rows[1].cells[1].paragraphs[0].add_run("Key Library: ").bold = True
    info_table.rows[1].cells[1].paragraphs[0].add_run("igraph (R)")
    
    # Spacing after info box
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # ----------------------------------------------------
    # SECTION 1: OVERVIEW
    # ----------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    h1_run = h1.add_run("1. Objective and Overview")
    h1_run.font.size = Pt(14)
    h1_run.font.bold = True
    h1_run.font.color.rgb = PRIMARY_COLOR
    
    overview_p = doc.add_paragraph()
    overview_p.paragraph_format.space_after = Pt(8)
    overview_p.paragraph_format.line_spacing = 1.15
    overview_p.add_run(
        "Social Network Analysis (SNA) models real-world interaction structures as mathematical graphs consisting "
        "of vertices (nodes) representing entities/actors and edges representing directional or bidirectional relationships. "
        "This experiment implements an end-to-end network analysis workflow in R using the "
    )
    overview_p.add_run("igraph").bold = True
    overview_p.add_run(
        " package. Starting from the dataset ("
    )
    overview_p.add_run("networkdata.csv").italic = True
    overview_p.add_run(
        "), the network was constructed, evaluated using fundamental network measures (degree, diameter, density, "
        "reciprocity, betweenness), visualized under multiple force-directed and geometric layouts, evaluated for "
        "Hub and Authority prominence (Kleinberg's HITS algorithm), and segmented into distinct community clusters using "
        "edge-betweenness modularity."
    )
    
    # ----------------------------------------------------
    # SECTION 2: PLOT 1 - INITIAL NETWORK GRAPH
    # ----------------------------------------------------
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    h2_run = h2.add_run("2. Initial Social Network Graph (Directed)")
    h2_run.font.size = Pt(14)
    h2_run.font.bold = True
    h2_run.font.color.rgb = PRIMARY_COLOR
    
    p2_desc = doc.add_paragraph()
    p2_desc.paragraph_format.space_after = Pt(6)
    p2_desc.paragraph_format.line_spacing = 1.15
    p2_desc.add_run(
        "The directed network was instantiated using "
    )
    p2_desc.add_run("graph.data.frame(df, directed = TRUE)").font.name = 'Consolas'
    p2_desc.add_run(
        ". In this initial visualization, all vertices are displayed with a uniform size ("
    )
    p2_desc.add_run("vertex.size = 12").font.name = 'Consolas'
    p2_desc.add_run(
        ") and colored green, with directional edge arrows depicting the flow of interaction. This allows initial "
        "inspection of overall graph density, identifying a highly interconnected central core surrounded by peripheral radiating actors."
    )
    
    # Add Image 1
    img1_path = "Social_Graph_Plot_1.png"
    if os.path.exists(img1_path):
        img_p1 = doc.add_paragraph()
        img_p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p1.paragraph_format.space_before = Pt(4)
        img_p1.paragraph_format.space_after = Pt(2)
        run1 = img_p1.add_run()
        run1.add_picture(img1_path, width=Inches(4.6))
        
        cap1 = doc.add_paragraph()
        cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap1.paragraph_format.space_after = Pt(14)
        crun1 = cap1.add_run("Figure 1: Initial Directed Social Network Graph with Uniform Node Sizing (igraph)")
        crun1.font.size = Pt(9.5)
        crun1.font.italic = True
        crun1.font.color.rgb = MUTED_COLOR

    # Page Break for next section
    doc.add_page_break()

    # ----------------------------------------------------
    # SECTION 3: PLOT 2 - FRUCHTERMAN-REINGOLD LAYOUT
    # ----------------------------------------------------
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3_run = h3.add_run("3. Degree-Weighted Visualization (Fruchterman-Reingold Layout)")
    h3_run.font.size = Pt(14)
    h3_run.font.bold = True
    h3_run.font.color.rgb = PRIMARY_COLOR
    
    p3_desc = doc.add_paragraph()
    p3_desc.paragraph_format.space_after = Pt(6)
    p3_desc.paragraph_format.line_spacing = 1.15
    p3_desc.add_run(
        "To highlight structural centrality and prominence, each node's size was scaled proportionally to its total degree ("
    )
    p3_desc.add_run("vertex.size = V(net)$degree * 0.4").font.name = 'Consolas'
    p3_desc.add_run(
        ") and assigned distinct colors from the "
    )
    p3_desc.add_run("rainbow(52)").font.name = 'Consolas'
    p3_desc.add_run(
        " palette. The graph was organized using the force-directed "
    )
    p3_desc.add_run("Fruchterman-Reingold algorithm").bold = True
    p3_desc.add_run(
        " ("
    )
    p3_desc.add_run("layout_with_fr").font.name = 'Consolas'
    p3_desc.add_run(
        "), which simulates repulsive electrical forces between nodes and attractive spring forces along edges, "
        "pulling heavily connected nodes (such as CA, CC, CD, and DD) into prominent central positions."
    )
    
    # Add Image 2
    img2_path = "Social_Graph_Plot_2.png"
    if os.path.exists(img2_path):
        img_p2 = doc.add_paragraph()
        img_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p2.paragraph_format.space_before = Pt(4)
        img_p2.paragraph_format.space_after = Pt(2)
        run2 = img_p2.add_run()
        run2.add_picture(img2_path, width=Inches(4.6))
        
        cap2 = doc.add_paragraph()
        cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap2.paragraph_format.space_after = Pt(14)
        crun2 = cap2.add_run("Figure 2: Fruchterman-Reingold Force-Directed Network with Degree-Proportional Node Sizing")
        crun2.font.size = Pt(9.5)
        crun2.font.italic = True
        crun2.font.color.rgb = MUTED_COLOR

    # ----------------------------------------------------
    # SECTION 4: PLOT 3 - KAMADA-KAWAI LAYOUT
    # ----------------------------------------------------
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(4)
    h4_run = h4.add_run("4. Degree-Weighted Visualization (Kamada-Kawai Layout)")
    h4_run.font.size = Pt(14)
    h4_run.font.bold = True
    h4_run.font.color.rgb = PRIMARY_COLOR
    
    p4_desc = doc.add_paragraph()
    p4_desc.paragraph_format.space_after = Pt(6)
    p4_desc.paragraph_format.line_spacing = 1.15
    p4_desc.add_run(
        "The graph was alternatively rendered using the "
    )
    p4_desc.add_run("Kamada-Kawai algorithm").bold = True
    p4_desc.add_run(
        " ("
    )
    p4_desc.add_run("layout_with_kk").font.name = 'Consolas'
    p4_desc.add_run(
        "). While Fruchterman-Reingold uses a force-directed model, Kamada-Kawai formulates the layout as an energy-minimization "
        "problem where Euclidean distances between pairs of nodes approximate their graph-theoretic shortest-path distances. "
        "This layout accentuates geometric symmetry, clearly grouping closely coupled actors and revealing the relative path distance of peripheral nodes."
    )
    
    # Add Image 3
    img3_path = "Social_Graph_Plot_3.png"
    if os.path.exists(img3_path):
        img_p3 = doc.add_paragraph()
        img_p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p3.paragraph_format.space_before = Pt(4)
        img_p3.paragraph_format.space_after = Pt(2)
        run3 = img_p3.add_run()
        run3.add_picture(img3_path, width=Inches(4.6))
        
        cap3 = doc.add_paragraph()
        cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap3.paragraph_format.space_after = Pt(14)
        crun3 = cap3.add_run("Figure 3: Kamada-Kawai Energy-Minimization Network Layout with Degree Scaling")
        crun3.font.size = Pt(9.5)
        crun3.font.italic = True
        crun3.font.color.rgb = MUTED_COLOR

    # Page Break for next section
    doc.add_page_break()

    # ----------------------------------------------------
    # SECTION 5: PLOT 4 - HUBS & AUTHORITIES
    # ----------------------------------------------------
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(8)
    h5.paragraph_format.space_after = Pt(4)
    h5_run = h5.add_run("5. Hubs and Authorities Analysis (Kleinberg's HITS Algorithm)")
    h5_run.font.size = Pt(14)
    h5_run.font.bold = True
    h5_run.font.color.rgb = PRIMARY_COLOR
    
    p5_desc = doc.add_paragraph()
    p5_desc.paragraph_format.space_after = Pt(6)
    p5_desc.paragraph_format.line_spacing = 1.15
    p5_desc.add_run(
        "In directed networks, centrality can be bifurcated into two mutually reinforcing roles using Jon Kleinberg's "
    )
    p5_desc.add_run("Hyperlink-Induced Topic Search (HITS)").bold = True
    p5_desc.add_run(" algorithm:\n")
    
    p5_bullet1 = doc.add_paragraph(style='List Bullet')
    p5_bullet1.paragraph_format.space_after = Pt(2)
    p5_bullet1.paragraph_format.line_spacing = 1.15
    r_b1_title = p5_bullet1.add_run("Hubs (hub_score): ")
    r_b1_title.bold = True
    p5_bullet1.add_run(
        "Nodes that point to many authoritative sources of information. In the visualization, nodes such as "
    )
    p5_bullet1.add_run("CC").bold = True
    p5_bullet1.add_run(" and ")
    p5_bullet1.add_run("CB").bold = True
    p5_bullet1.add_run(" exhibit large vertex radii, signifying high hub status.")
    
    p5_bullet2 = doc.add_paragraph(style='List Bullet')
    p5_bullet2.paragraph_format.space_after = Pt(6)
    p5_bullet2.paragraph_format.line_spacing = 1.15
    r_b2_title = p5_bullet2.add_run("Authorities (authority.score): ")
    r_b2_title.bold = True
    p5_bullet2.add_run(
        "Nodes that receive inward links from multiple reliable hubs. Node "
    )
    p5_bullet2.add_run("CA").bold = True
    p5_bullet2.add_run(
        " clearly emerges as the preeminent authority in this network, receiving widespread endorsement from peer hubs."
    )
    
    # Add Image 4 (Wide plot)
    img4_path = "Social_Graph_Plot_Hubs_&_Authorities.png"
    if os.path.exists(img4_path):
        img_p4 = doc.add_paragraph()
        img_p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p4.paragraph_format.space_before = Pt(4)
        img_p4.paragraph_format.space_after = Pt(2)
        run4 = img_p4.add_run()
        run4.add_picture(img4_path, width=Inches(6.2))
        
        cap4 = doc.add_paragraph()
        cap4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap4.paragraph_format.space_after = Pt(14)
        crun4 = cap4.add_run("Figure 4: Comparative Layout of Hubs (left) and Authorities (right) under HITS Analysis")
        crun4.font.size = Pt(9.5)
        crun4.font.italic = True
        crun4.font.color.rgb = MUTED_COLOR

    # Page Break for next section
    doc.add_page_break()

    # ----------------------------------------------------
    # SECTION 6: PLOT 5 - COMMUNITY DETECTION
    # ----------------------------------------------------
    h6 = doc.add_paragraph()
    h6.paragraph_format.space_before = Pt(8)
    h6.paragraph_format.space_after = Pt(4)
    h6_run = h6.add_run("6. Community Detection & Subgroup Clustering")
    h6_run.font.size = Pt(14)
    h6_run.font.bold = True
    h6_run.font.color.rgb = PRIMARY_COLOR
    
    p6_desc = doc.add_paragraph()
    p6_desc.paragraph_format.space_after = Pt(6)
    p6_desc.paragraph_format.line_spacing = 1.15
    p6_desc.add_run(
        "Community detection identifies modular clusters and cohesive subgroups within the interaction graph. "
        "The graph was converted to an undirected representation and analyzed using the "
    )
    p6_desc.add_run("Girvan-Newman edge-betweenness community detection").bold = True
    p6_desc.add_run(
        " method ("
    )
    p6_desc.add_run("cluster_edge_betweenness").font.name = 'Consolas'
    p6_desc.add_run(
        "). This algorithm progressively removes edges with the highest betweenness centrality (which bridge distinct communities), "
        "uncovering natural organizational clusters. Shaded convex hulls delineate distinct functional subgroups within the population."
    )
    
    # Add Image 5 (Wide plot)
    img5_path = "group_in_networks.jpeg"
    if os.path.exists(img5_path):
        img_p6 = doc.add_paragraph()
        img_p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p6.paragraph_format.space_before = Pt(4)
        img_p6.paragraph_format.space_after = Pt(2)
        run6 = img_p6.add_run()
        run6.add_picture(img5_path, width=Inches(6.2))
        
        cap6 = doc.add_paragraph()
        cap6.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap6.paragraph_format.space_after = Pt(14)
        crun6 = cap6.add_run("Figure 5: Network Community Detection via Edge-Betweenness Clustering with Subgroup Outlines")
        crun6.font.size = Pt(9.5)
        crun6.font.italic = True
        crun6.font.color.rgb = MUTED_COLOR

    # ----------------------------------------------------
    # SECTION 7: SUMMARY & NETWORK METRICS
    # ----------------------------------------------------
    h7 = doc.add_paragraph()
    h7.paragraph_format.space_before = Pt(10)
    h7.paragraph_format.space_after = Pt(4)
    h7_run = h7.add_run("7. Summary of Network Characteristics & Findings")
    h7_run.font.size = Pt(14)
    h7_run.font.bold = True
    h7_run.font.color.rgb = PRIMARY_COLOR
    
    # Summary Table
    table = doc.add_table(rows=7, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    
    t_widths = [Inches(2.5), Inches(4.2)]
    
    headers = [
        ("Network Metric / Parameter", "Value / Observation"),
        ("Total Vertices (Nodes)", "52 actors"),
        ("Total Directed Edges", "290 interactions"),
        ("Average Node Degree", "11.15 interactions per node"),
        ("Most Influential Node (Degree)", "Node 'CA' (highest overall degree of 62)"),
        ("Primary Hubs & Authorities", "Hubs: CC, CB | Authorities: CA, CD, DD"),
        ("Community Structure", "Multiple dense clusters identified via Girvan-Newman edge betweenness")
    ]
    
    for row_idx, (m_lbl, m_val) in enumerate(headers):
        row = table.rows[row_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = t_widths[0], t_widths[1]
        set_cell_margins(c0, top=70, bottom=70, left=100, right=100)
        set_cell_margins(c1, top=70, bottom=70, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        p1 = c1.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        
        if row_idx == 0:
            set_cell_background(c0, "1B365D")
            set_cell_background(c1, "1B365D")
            r0 = p0.add_run(m_lbl)
            r1 = p1.add_run(m_val)
            r0.bold, r1.bold = True, True
            r0.font.color.rgb = RGBColor(255, 255, 255)
            r1.font.color.rgb = RGBColor(255, 255, 255)
        else:
            bg_color = "F9FBFD" if row_idx % 2 == 1 else "FFFFFF"
            set_cell_background(c0, bg_color)
            set_cell_background(c1, bg_color)
            r0 = p0.add_run(m_lbl)
            r1 = p1.add_run(m_val)
            r0.bold = True
            r0.font.size = Pt(10)
            r1.font.size = Pt(10)

    # Key Takeaways paragraph
    concl_p = doc.add_paragraph()
    concl_p.paragraph_format.space_before = Pt(10)
    concl_p.paragraph_format.space_after = Pt(0)
    concl_p.paragraph_format.line_spacing = 1.15
    concl_p.add_run(
        "Conclusion: The social network demonstrates a core-periphery architecture typical of real-world collaborative networks. "
        "The analysis demonstrates how vertex sizing, layout algorithms, HITS score decomposition, and edge-betweenness community clustering "
        "provide deep insight into the structure, influential members, and organizational groupings within social systems."
    )

    output_filename = "Assignment_5_Social_Network_Analysis.docx"
    doc.save(output_filename)
    print(f"Document successfully created and saved as: {output_filename}")

if __name__ == "__main__":
    build_word_document()
